"""
키워드 연구 보고서 생성

첨부된 네이버 키워드 연구 XLSX를 읽어 GitHub Pages용 날짜별 보고서를 만든다.
원본 XLSX는 읽기 전용으로 사용하며, 매일 JSON/CSV 스냅샷을 남긴다.
검색광고 API 연동 전에도 기준 파일을 이용해 보고서를 확인할 수 있도록 설계했다.
"""
import csv
import datetime as dt
import base64
import hashlib
import hmac
import html
import json
import os
import re
import shutil
import time
from collections import Counter, defaultdict

from openpyxl import load_workbook

try:
    import requests
except ImportError:  # 로컬 점검용 런타임에 requests가 없어도 엑셀 기준 보고서는 생성한다.
    requests = None

KST = dt.timezone(dt.timedelta(hours=9))
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
DOCS = os.path.join(ROOT, "docs")
REPORT_DATA = os.path.join(DATA, "keyword_reports")
KEYWORD_DOCS = os.path.join(DOCS, "keywords")
DOWNLOADS = os.path.join(KEYWORD_DOCS, "downloads")
SOURCE_XLSX = os.path.join(DATA, "naver_keyword_research_1000_2026-09-04_fixed.xlsx")
SEARCHAD_URL = "https://api.searchad.naver.com"
SEARCHAD_PATH = "/keywordstool"
SEARCHAD_BATCH_SIZE = 5

WEEKDAY = ["월", "화", "수", "목", "금", "토", "일"]


def esc(value):
    return html.escape(str(value or ""), quote=True)


def number(value):
    if isinstance(value, bool):
        return None
    try:
        return float(value) if value is not None and value != "" else None
    except (TypeError, ValueError):
        return None


def integer(value):
    value = number(value)
    return int(round(value)) if value is not None else None


def display_number(value):
    value = number(value)
    if value is None:
        return "-"
    if value == int(value):
        return f"{int(value):,}"
    return f"{value:,.1f}"


def display_score(value):
    value = number(value)
    return "-" if value is None else f"{value:.1f}"


def clean(value):
    return str(value).strip() if value is not None else ""


def searchad_credentials():
    return {
        "access_license": os.getenv("NAVER_SEARCHAD_ACCESS_LICENSE", "").strip(),
        "secret_key": os.getenv("NAVER_SEARCHAD_SECRET_KEY", "").strip(),
        "customer_id": os.getenv("NAVER_SEARCHAD_CUSTOMER_ID", "").strip(),
    }


def normalize_keyword(value):
    return re.sub(r"\s+", "", clean(value)).lower()


def searchad_signature(timestamp, method, uri, secret_key):
    message = f"{timestamp}.{method}.{uri}"
    digest = hmac.new(
        secret_key.encode("utf-8"), message.encode("utf-8"), hashlib.sha256
    ).digest()
    return base64.b64encode(digest).decode("ascii")


def parse_search_count(value):
    text = clean(value).replace(",", "")
    if not text:
        return 0, 0
    if text.startswith("<"):
        try:
            upper = max(0, int(float(text[1:].strip())) - 1)
        except ValueError:
            upper = 0
        return 0, upper
    try:
        exact = int(float(text))
    except ValueError:
        return 0, 0
    return exact, exact


def searchad_headers(credentials):
    timestamp = str(int(time.time() * 1000))
    return {
        "Content-Type": "application/json; charset=UTF-8",
        "X-Timestamp": timestamp,
        "X-API-KEY": credentials["access_license"],
        "X-Customer": credentials["customer_id"],
        "X-Signature": searchad_signature(
            timestamp, "GET", SEARCHAD_PATH, credentials["secret_key"]
        ),
    }


def fetch_searchad_batch(keywords, credentials):
    params = {
        "hintKeywords": ",".join(keywords),
        "includeHintKeywords": "1",
        "showDetail": "1",
    }
    for attempt in range(3):
        try:
            response = requests.get(
                f"{SEARCHAD_URL}{SEARCHAD_PATH}",
                params=params,
                headers=searchad_headers(credentials),
                timeout=20,
            )
            if response.status_code == 200:
                body = response.json()
                result = {}
                for item in body.get("keywordList", []):
                    key = normalize_keyword(item.get("relKeyword"))
                    if key:
                        result[key] = item
                return result, ""
            if response.status_code == 429 or response.status_code >= 500:
                if attempt < 2:
                    time.sleep(2 ** attempt)
                    continue
            detail = response.text[:240].replace("\n", " ")
            return {}, f"HTTP {response.status_code}: {detail}"
        except (requests.RequestException, ValueError) as exc:
            if attempt < 2:
                time.sleep(2 ** attempt)
                continue
            return {}, f"{type(exc).__name__}: {exc}"
    return {}, "검색광고 API 요청 실패"


def refresh_with_searchad(keyword_rows):
    credentials = searchad_credentials()
    missing = [key for key, value in credentials.items() if not value]
    if missing:
        return {
            "enabled": False,
            "reason": "검색광고 API Secret 3개가 모두 등록되지 않았습니다.",
            "missing": missing,
            "checked_at": "",
            "updated_count": 0,
            "failed_count": 0,
        }
    if requests is None:
        return {
            "enabled": False,
            "reason": "requests 패키지가 설치되지 않아 검색광고 API를 건너뛰었습니다.",
            "missing": [],
            "checked_at": "",
            "updated_count": 0,
            "failed_count": 0,
        }

    checked_at = dt.datetime.now(KST).isoformat(timespec="seconds")
    updated = 0
    failed = 0
    errors = []
    for start in range(0, len(keyword_rows), SEARCHAD_BATCH_SIZE):
        batch = keyword_rows[start:start + SEARCHAD_BATCH_SIZE]
        requested = [row["keyword"] for row in batch if row["keyword"]]
        if not requested:
            continue
        values, error = fetch_searchad_batch(requested, credentials)
        if error and len(requested) > 1:
            # 여러 키워드 요청이 계정/API 제한에 걸리면 해당 묶음만 단건으로 재시도한다.
            values = {}
            error = ""
            for keyword in requested:
                single, single_error = fetch_searchad_batch([keyword], credentials)
                values.update(single)
                if single_error:
                    errors.append(single_error)
                time.sleep(0.08)
        elif error:
            errors.append(error)

        for row in batch:
            item = values.get(normalize_keyword(row["keyword"]))
            if not item:
                failed += 1
                continue
            pc_lower, pc_upper = parse_search_count(item.get("monthlyPcQcCnt"))
            mobile_lower, mobile_upper = parse_search_count(item.get("monthlyMobileQcCnt"))
            row["source_monthly_lower"] = row.get("monthly_lower")
            row["source_monthly_upper"] = row.get("monthly_upper")
            row["monthly_lower"] = pc_lower + mobile_lower
            row["monthly_upper"] = pc_upper + mobile_upper
            row["live_pc_monthly"] = pc_lower
            row["live_mobile_monthly"] = mobile_lower
            row["live_competition"] = clean(item.get("compIdx"))
            row["live_depth"] = clean(item.get("plAvgDepth"))
            row["live_checked_at"] = checked_at
            row["search_source"] = "네이버 검색광고 API"
            updated += 1
        time.sleep(0.08)

    return {
        "enabled": True,
        "reason": "",
        "missing": [],
        "checked_at": checked_at,
        "updated_count": updated,
        "failed_count": failed,
        "error_count": len(errors),
        "errors": errors[:3],
    }


def sheet_rows(workbook, sheet_name, header_row=4):
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [clean(value) for value in rows[header_row - 1]]
    output = []
    for row_number, values in enumerate(rows[header_row:], header_row + 1):
        if not any(value not in (None, "") for value in values):
            continue
        item = {headers[index]: values[index] if index < len(values) else None
                for index in range(len(headers)) if headers[index]}
        item["_row"] = row_number
        output.append(item)
    return output


def find_as_of(workbook):
    for sheet_name in ("Data & Targets", "기회점수", "TOP30"):
        if sheet_name not in workbook.sheetnames:
            continue
        for row in workbook[sheet_name].iter_rows(min_row=1, max_row=3, values_only=True):
            text = " ".join(clean(value) for value in row if value not in (None, ""))
            match = re.search(r"(20\d{2}-\d{2}-\d{2})", text)
            if match:
                return match.group(1)
    match = re.search(r"(20\d{2}-\d{2}-\d{2})", os.path.basename(SOURCE_XLSX))
    return match.group(1) if match else "기준일 미상"


def as_text_date(value):
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dt.date):
        return value.isoformat()
    return clean(value)


def build_report(today):
    if not os.path.exists(SOURCE_XLSX):
        raise FileNotFoundError(f"키워드 기준 파일이 없습니다: {SOURCE_XLSX}")

    workbook = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    targets = sheet_rows(workbook, "Data & Targets")
    scores = sheet_rows(workbook, "기회점수")
    top30 = sheet_rows(workbook, "TOP30")
    raw100 = sheet_rows(workbook, "원본100")
    alternatives = sheet_rows(workbook, "대체어40")

    score_by_id = {clean(row.get("ID")): row for row in scores if clean(row.get("ID"))}
    keyword_rows = []
    for row in targets:
        key = clean(row.get("ID"))
        score = score_by_id.get(key, {})
        keyword_rows.append({
            "id": key,
            "category": clean(row.get("주제")),
            "category_rank": integer(row.get("주제내 검색순위")),
            "keyword": clean(row.get("검색어(조회 표기)")),
            "search_type": clean(row.get("검색 유형")),
            "monthly_lower": integer(row.get("월검색 합계 하한")),
            "monthly_upper": integer(row.get("월검색 합계 상한")),
            "blog_documents": integer(row.get("누적 블로그 문서")),
            "recent_posts": integer(row.get("최근30일 발행수")),
            "recent_posts_note": clean(row.get("발행량 집계 범위")),
            "search_checked_at": as_text_date(row.get("검색량 조회일(KST)")),
            "opportunity": number(score.get("기회점수")),
            "persistence": number(score.get("지속성 근거점수")),
            "home_fit": number(score.get("홈판 제목화 가능성")),
            "priority": number(score.get("발행 우선점수")),
            "overall_rank": integer(score.get("전체 발행순위")),
            "search_source": "첨부 엑셀",
            "live_checked_at": "",
            "live_pc_monthly": None,
            "live_mobile_monthly": None,
            "live_competition": "",
            "live_depth": "",
        })

    searchad = refresh_with_searchad(keyword_rows)

    categories = []
    grouped = defaultdict(list)
    for row in keyword_rows:
        grouped[row["category"]].append(row)
    for category, rows in grouped.items():
        priority_rows = sorted(
            rows, key=lambda item: (item["priority"] is not None,
                                    item["priority"] or -1), reverse=True
        )
        volume_rows = sorted(rows, key=lambda item: item["monthly_lower"] or 0, reverse=True)
        values = [row["monthly_lower"] for row in rows if row["monthly_lower"] is not None]
        documents = [row["blog_documents"] for row in rows if row["blog_documents"] is not None]
        categories.append({
            "category": category,
            "count": len(rows),
            "median_monthly_lower": sorted(values)[(len(values) - 1) // 2] if values else None,
            "median_documents": sorted(documents)[(len(documents) - 1) // 2] if documents else None,
            "top_keyword": volume_rows[0]["keyword"] if volume_rows else "",
            "top_priority_keyword": priority_rows[0]["keyword"] if priority_rows else "",
            "top_priority": priority_rows[0]["priority"] if priority_rows else None,
            "top_rows": [
                {
                    "id": item["id"],
                    "keyword": item["keyword"],
                    "monthly_lower": item["monthly_lower"],
                    "priority": item["priority"],
                    "opportunity": item["opportunity"],
                }
                for item in priority_rows[:5]
            ],
        })
    categories.sort(key=lambda item: item["category"])

    report_top30 = []
    keyword_by_id = {row["id"]: row for row in keyword_rows if row["id"]}
    for row in top30:
        key = clean(row.get("ID"))
        live = keyword_by_id.get(key, {})
        report_top30.append({
            "order": integer(row.get("선별순서")),
            "id": key,
            "category": clean(row.get("주제")),
            "keyword": clean(row.get("검색어")),
            "monthly_lower": live.get("monthly_lower") if live.get("live_checked_at") else integer(row.get("월검색수 하한")),
            "opportunity": number(row.get("기회점수")),
            "persistence": number(row.get("지속성 근거점수")),
            "home_fit": number(row.get("홈판 제목화 가능성")),
            "priority": number(row.get("발행 우선점수")),
            "overall_rank": integer(row.get("현재 전체순위")),
            "explanation": clean(row.get("선별 설명")),
        })

    report = {
        "date": today,
        "generated_at": dt.datetime.now(KST).isoformat(),
        "source_file": os.path.basename(SOURCE_XLSX),
        "source_as_of": find_as_of(workbook),
        "summary": {
            "candidate_count": len(keyword_rows),
            "category_count": len(categories),
            "top30_count": len(report_top30),
            "raw100_count": len(raw100),
            "alternative_count": len(alternatives),
            "maintained_raw_count": sum(1 for row in raw100 if clean(row.get("상태")) == "유지 가능"),
            "alternative_review_count": sum(1 for row in raw100 if clean(row.get("상태")) == "대체어 검토"),
            "searchad_api_enabled": searchad["enabled"],
            "searchad_live_count": searchad["updated_count"],
        },
        "searchad": searchad,
        "categories": categories,
        "top30": report_top30,
        "keywords": keyword_rows,
        "alternatives": [
            {
                "category": clean(row.get("주제")),
                "original": clean(row.get("원본 검색어")),
                "recommended": clean(row.get("추천 검색어")),
                "relationship": clean(row.get("의도 관계")),
                "reason": clean(row.get("추천 근거와 차이")),
                "volume": integer(row.get("대체 월검색 하한")),
            }
            for row in alternatives
        ],
    }
    workbook.close()
    return report


def write_json(report):
    os.makedirs(REPORT_DATA, exist_ok=True)
    path = os.path.join(REPORT_DATA, f"keyword-{report['date']}.json")
    with open(path, "w", encoding="utf-8") as file:
        json.dump(report, file, ensure_ascii=False, indent=2)
    return path


def write_csv(report):
    os.makedirs(DOWNLOADS, exist_ok=True)
    path = os.path.join(DOWNLOADS, f"keyword-report-{report['date']}.csv")
    fields = [
        "id", "category", "category_rank", "keyword", "search_type",
        "monthly_lower", "monthly_upper", "blog_documents", "recent_posts",
        "recent_posts_note", "opportunity", "persistence", "home_fit",
        "priority", "overall_rank", "search_source", "live_pc_monthly",
        "live_mobile_monthly", "live_competition", "live_depth", "live_checked_at",
    ]
    with open(path, "w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerows({field: row.get(field, "") for field in fields} for row in report["keywords"])
    return path


def report_date_label(value):
    try:
        date = dt.datetime.strptime(value, "%Y-%m-%d")
        return f"{date.year}년 {date.month}월 {date.day}일 ({WEEKDAY[date.weekday()]})"
    except ValueError:
        return value


def render_css():
    return """
*{box-sizing:border-box}body{margin:0;background:#f7f6f3;color:#1a1a1a;font-family:-apple-system,BlinkMacSystemFont,"Apple SD Gothic Neo","Malgun Gothic",sans-serif;line-height:1.55}.wrap{max-width:1200px;margin:0 auto;padding:22px 16px 60px}header{border-bottom:1px solid #e3e0da;padding:12px 0 20px;margin-bottom:20px}h1{font-size:1.5rem;margin:0 0 4px;letter-spacing:-.03em}.date,.muted{color:#6b6b6b;font-size:.82rem}.nav{display:flex;flex-wrap:wrap;gap:8px;margin-top:14px}.nav a,.button{border:1px solid #e3e0da;background:#fff;color:#555;border-radius:99px;padding:6px 11px;text-decoration:none;font-size:.78rem}.nav a:hover,.button:hover{color:#111}.notice{background:#fff7ed;border:1px solid #fed7aa;color:#9a3412;border-radius:10px;padding:12px 14px;margin:14px 0}.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin:16px 0}.stat{background:#fff;border:1px solid #e3e0da;border-radius:10px;padding:13px}.stat strong{display:block;font-size:1.35rem}.stat span{color:#6b6b6b;font-size:.76rem}.panel{background:#fff;border:1px solid #e3e0da;border-radius:12px;padding:15px;margin:18px 0}.panel h2{font-size:1.1rem;margin:0 0 10px}.tools{display:flex;gap:8px;flex-wrap:wrap;margin:10px 0}.tools input,.tools select{border:1px solid #d8d4cc;background:#fff;border-radius:7px;padding:8px 9px;font:inherit;font-size:.82rem}.scroll{overflow-x:auto}.table{width:100%;border-collapse:collapse;font-size:.78rem}.table th,.table td{border-bottom:1px solid #e3e0da;padding:8px;text-align:left;vertical-align:top}.table th{white-space:nowrap;color:#6b6b6b;font-weight:600}.table a{color:#1a1a1a;font-weight:600}.tag{display:inline-block;background:#efece6;border-radius:99px;color:#666;padding:2px 7px;font-size:.7rem;white-space:nowrap}.check{white-space:nowrap;color:#666}.check input{accent-color:#3b8f5b}.top-row{background:#fbfaf7}.top-row td:first-child{font-weight:700}.archive{list-style:none;padding:0;margin:0}.archive li{padding:10px 0;border-bottom:1px solid #e3e0da}.archive a{color:#1a1a1a;text-decoration:none}.small{font-size:.72rem;color:#777}footer{border-top:1px solid #e3e0da;margin-top:34px;padding-top:14px;color:#777;font-size:.76rem}@media(max-width:700px){.table{min-width:950px}.wrap{padding-left:10px;padding-right:10px}}
"""


def nav_html():
    return '<nav class="nav"><a href="../">오늘의 뉴스 브리핑</a><a href="./">키워드 연구</a><a href="archive.html">키워드 날짜별 보기</a></nav>'


def keyword_row_html(row, report_date):
    keyword_id = f"keyword:{report_date}:{row['id']}"
    competition = row.get("live_competition") or "-"
    checked_at = row.get("live_checked_at") or "-"
    return (
        f'<tr id="kw-{esc(row["id"])}" data-category="{esc(row["category"])}" '
        f'data-type="{esc(row["search_type"])}" data-search="{esc((row["category"] + " " + row["keyword"]).lower())}">'
        f'<td>{esc(row["id"])}</td><td>{esc(row["category"])}</td>'
        f'<td><strong>{esc(row["keyword"])}</strong><br><span class="tag">{esc(row["search_type"])}</span></td>'
        f'<td>{display_number(row["monthly_lower"])}</td><td>{esc(competition)}</td><td>{display_number(row["blog_documents"])}</td>'
        f'<td>{display_number(row["recent_posts"])}<br><span class="small">{esc(row["recent_posts_note"])}</span></td>'
        f'<td>{display_score(row["opportunity"])}</td><td>{display_score(row["persistence"])}</td>'
        f'<td>{display_score(row["home_fit"])}</td><td>{display_score(row["priority"])}</td><td class="small">{esc(checked_at)}</td>'
        f'<td class="check"><label><input class="keyword-check" type="checkbox" data-keyword-id="{esc(keyword_id)}"> 완료</label></td>'
        '</tr>'
    )


def render_report(report):
    categories = "".join(
        f'<tr><td>{esc(item["category"])}</td><td>{item["count"]}</td>'
        f'<td>{display_number(item["median_monthly_lower"])}</td><td>{display_number(item["median_documents"])}</td>'
        f'<td>{esc(item["top_keyword"])}</td><td>{esc(item["top_priority_keyword"])} '
        f'<span class="small">({display_score(item["top_priority"])})</span></td></tr>'
        for item in report["categories"]
    )
    keyword_by_id = {row["id"]: row for row in report["keywords"]}
    top30 = "".join(
        f'<tr class="top-row"><td>{item["order"]}</td><td>{esc(item["category"])}</td>'
        f'<td><a href="#kw-{esc(item["id"])}">{esc(item["keyword"])}</a></td>'
        f'<td>{display_number(item["monthly_lower"])}</td><td>{esc(keyword_by_id.get(item["id"], {}).get("live_competition") or "-")}</td><td>{display_score(item["opportunity"])}</td>'
        f'<td>{display_score(item["persistence"])}</td><td>{display_score(item["home_fit"])}</td>'
        f'<td>{display_score(item["priority"])}</td><td>{item["overall_rank"] or "-"}</td></tr>'
        for item in report["top30"]
    )
    keywords = "".join(keyword_row_html(row, report["date"]) for row in report["keywords"])
    alternatives = "".join(
        f'<tr><td>{esc(item["category"])}</td><td>{esc(item["original"])}</td>'
        f'<td>{esc(item["recommended"])}</td><td><span class="tag">{esc(item["relationship"])}</span></td>'
        f'<td>{display_number(item["volume"])}</td><td>{esc(item["reason"])}</td></tr>'
        for item in report["alternatives"]
    )
    categories_options = ''.join(
        f'<option value="{esc(item["category"])}">{esc(item["category"])}</option>'
        for item in report["categories"]
    )
    search_types = sorted({row["search_type"] for row in report["keywords"] if row["search_type"]})
    type_options = ''.join(f'<option value="{esc(value)}">{esc(value)}</option>' for value in search_types)
    summary = report["summary"]
    searchad = report.get("searchad", {})
    if searchad.get("enabled"):
        notice = (
            f'검색광고 API로 {searchad.get("updated_count", 0):,}개 키워드의 월간 검색량·경쟁도 일부를 갱신했습니다. '
            f'갱신 시각: {esc(searchad.get("checked_at", ""))}. 기회점수·발행 우선점수는 원본 연구 파일 기준입니다.'
        )
    else:
        notice = (
            '검색광고 API Secret이 모두 등록되면 월간 검색량·경쟁도를 자동 갱신합니다. '
            '현재 보고서는 첨부 엑셀의 조회 기준일 자료를 사용합니다.'
        )
    script = """
<script>
(() => {
  const storageKey = 'blogKeywordCompleted';
  const read = () => { try { return JSON.parse(localStorage.getItem(storageKey) || '{}'); } catch (_) { return {}; } };
  const write = value => { try { localStorage.setItem(storageKey, JSON.stringify(value)); } catch (_) {} };
  const refresh = () => {
    const saved = read(); let done = 0;
    document.querySelectorAll('.keyword-check').forEach(input => {
      const checked = Boolean(saved[input.dataset.keywordId]); input.checked = checked; if (checked) done += 1;
    });
    document.querySelectorAll('[data-summary]').forEach(el => el.textContent = `완료 ${done} / 전체 ${document.querySelectorAll('.keyword-check').length}`);
  };
  document.querySelectorAll('.keyword-check').forEach(input => input.addEventListener('change', () => {
    const saved = read(); if (input.checked) saved[input.dataset.keywordId] = true; else delete saved[input.dataset.keywordId]; write(saved); refresh();
  }));
  document.querySelector('[data-reset]')?.addEventListener('click', () => {
    if (!confirm('이 보고서의 완료 표시를 모두 지울까요?')) return;
    const saved = read(); document.querySelectorAll('.keyword-check').forEach(input => delete saved[input.dataset.keywordId]); write(saved); refresh();
  });
  const category = document.querySelector('#category-filter'); const type = document.querySelector('#type-filter'); const search = document.querySelector('#keyword-search');
  const rows = [...document.querySelectorAll('#keyword-rows tr[data-category]')];
  const apply = () => { const q = (search.value || '').trim().toLowerCase(); rows.forEach(row => { const ok = (!category.value || row.dataset.category === category.value) && (!type.value || row.dataset.type === type.value) && (!q || row.dataset.search.includes(q)); row.hidden = !ok; }); };
  category?.addEventListener('change', apply); type?.addEventListener('change', apply); search?.addEventListener('input', apply); refresh();
})();
</script>
"""
    return f'''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="keyword-report-date" content="{esc(report["date"])}"><title>{report_date_label(report["date"])} 키워드 연구 보고서</title><style>{render_css()}</style></head><body><div class="wrap">
<header><h1>오늘의 네이버 키워드 연구 보고서</h1><div class="date">{report_date_label(report["date"])} 생성 · 원본 조회 기준 {esc(report["source_as_of"])} · 첨부 엑셀 기반</div>{nav_html()}</header>
<div class="notice"><strong>자료 기준 안내</strong><br>{notice}<br>뉴스 브리핑과 별도의 두 번째 보고서입니다.</div>
<div class="stats"><div class="stat"><strong>{summary["candidate_count"]:,}</strong><span>전체 키워드</span></div><div class="stat"><strong>{summary["category_count"]}</strong><span>대분류</span></div><div class="stat"><strong>{summary["top30_count"]}</strong><span>TOP30</span></div><div class="stat"><strong>{summary["raw100_count"]}</strong><span>원본 키워드</span></div><div class="stat"><strong>{summary["alternative_count"]}</strong><span>대체어 검토</span></div><div class="stat"><strong>{summary.get("searchad_live_count", 0):,}</strong><span>API 갱신 키워드</span></div></div>
<section class="panel"><h2>오늘 먼저 볼 TOP30</h2><p class="muted">검색량·기회점수·지속성·홈판 적합도를 조합한 원본 파일의 우선순위입니다. 검색광고 API 연결 시 월검색 하한과 경쟁도가 최신 값으로 표시됩니다.</p><div class="scroll"><table class="table"><thead><tr><th>순위</th><th>카테고리</th><th>키워드</th><th>월검색 하한</th><th>경쟁도</th><th>기회</th><th>지속성</th><th>홈판</th><th>발행 우선</th><th>전체 순위</th></tr></thead><tbody>{top30}</tbody></table></div></section>
<section class="panel"><h2>카테고리별 요약</h2><div class="scroll"><table class="table"><thead><tr><th>카테고리</th><th>키워드 수</th><th>월검색 중앙값</th><th>누적 문서 중앙값</th><th>검색량 1위</th><th>발행 우선 1위</th></tr></thead><tbody>{categories}</tbody></table></div></section>
<section class="panel"><h2>전체 키워드 1,000개</h2><div class="tools"><select id="category-filter"><option value="">전체 카테고리</option>{categories_options}</select><select id="type-filter"><option value="">전체 검색 유형</option>{type_options}</select><input id="keyword-search" type="search" placeholder="키워드 검색"><span class="button" data-summary>완료 0 / 전체 0</span><button class="button" type="button" data-reset>완료 표시 초기화</button></div><p class="muted">체크 상태는 현재 브라우저에 저장됩니다. 장기 발행 관리는 기존 작성 관리표를 함께 사용하세요.</p><div class="scroll"><table class="table"><thead><tr><th>ID</th><th>카테고리</th><th>키워드</th><th>월검색 하한</th><th>경쟁도</th><th>누적 문서</th><th>최근30일 발행</th><th>기회</th><th>지속성</th><th>홈판</th><th>발행 우선</th><th>API 갱신</th><th>완료</th></tr></thead><tbody id="keyword-rows">{keywords}</tbody></table></div></section>
<section class="panel"><h2>저검색 원본의 대체·확장안</h2><div class="scroll"><table class="table"><thead><tr><th>카테고리</th><th>원본</th><th>추천 검색어</th><th>관계</th><th>월검색 하한</th><th>차이·근거</th></tr></thead><tbody>{alternatives}</tbody></table></div></section>
<footer><a class="button" href="downloads/keyword-report-{esc(report["date"])}.csv">오늘 CSV 내려받기</a> <a class="button" href="downloads/source-keyword-research.xlsx">원본 엑셀 내려받기</a><br><br>이 보고서는 키워드 수요 연구 자료입니다. 최신 뉴스·정책·건강 정보의 사실 확인을 대신하지 않습니다.</footer>
</div>{script}</body></html>'''


def render_archive(reports):
    items = ''.join(
        f'<li><a href="{esc(report["date"])}.html">{report_date_label(report["date"])}</a>'
        f' <span class="small">· {report["summary"]["candidate_count"]:,}개 키워드 · 기준 {esc(report["source_as_of"])}</span></li>'
        for report in reports
    )
    return f'''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>키워드 연구 날짜별 보고서</title><style>{render_css()}</style></head><body><div class="wrap"><header><h1>키워드 연구 날짜별 보고서</h1><div class="date">첨부 키워드 연구 파일 기반</div>{nav_html()}</header><section class="panel"><h2>날짜별 보고서</h2><ul class="archive">{items or "<li>아직 보고서가 없습니다.</li>"}</ul></section><footer>뉴스 브리핑과 별도로 매일 생성되는 키워드 연구 보고서입니다.</footer></div></body></html>'''


def main():
    today = dt.datetime.now(KST).strftime("%Y-%m-%d")
    report = build_report(today)
    write_json(report)
    write_csv(report)
    os.makedirs(DOWNLOADS, exist_ok=True)
    shutil.copy2(SOURCE_XLSX, os.path.join(DOWNLOADS, "source-keyword-research.xlsx"))
    os.makedirs(KEYWORD_DOCS, exist_ok=True)
    with open(os.path.join(KEYWORD_DOCS, f"{today}.html"), "w", encoding="utf-8") as file:
        file.write(render_report(report))

    reports = []
    for path in sorted(
        (os.path.join(REPORT_DATA, name) for name in os.listdir(REPORT_DATA)
         if name.startswith("keyword-") and name.endswith(".json")), reverse=True
    ):
        try:
            with open(path, encoding="utf-8") as file:
                reports.append(json.load(file))
        except (OSError, json.JSONDecodeError):
            continue
    with open(os.path.join(KEYWORD_DOCS, "index.html"), "w", encoding="utf-8") as file:
        file.write(render_archive(reports))
    with open(os.path.join(KEYWORD_DOCS, "archive.html"), "w", encoding="utf-8") as file:
        file.write(render_archive(reports))
    print(f"키워드 연구 보고서 생성 완료 - {today} · {len(report['keywords'])}개 키워드")


if __name__ == "__main__":
    main()
